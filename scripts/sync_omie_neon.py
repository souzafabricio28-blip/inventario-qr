"""Sync Omie Excel -> Neon/Postgres em lote (rápido)."""
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from dotenv import load_dotenv
load_dotenv(ROOT / ".env")

import psycopg2
from psycopg2.extras import execute_batch, RealDictCursor
from openpyxl import load_workbook


def norm(val):
    if val is None:
        return ""
    c = str(val).strip()
    if c.endswith(".0") and c[:-2].replace("-", "").replace(".", "").isdigit():
        c = c[:-2]
    return c


def main():
    url = os.environ.get("DATABASE_URL", "").replace("&channel_binding=require", "")
    if not url:
        print("DATABASE_URL ausente")
        sys.exit(1)

    xlsx = ROOT / "omie_produtos.xlsx"
    if not xlsx.exists():
        xlsx = Path.home() / "Documents" / "lista de produtos cadastrados no omie.xlsx"
    if not xlsx.exists():
        print("Excel Omie nao encontrado")
        sys.exit(1)

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    linhas = list(wb.active.iter_rows(values_only=True))[1:]
    wb.close()

    omie = []
    for row in linhas:
        if not row:
            continue
        codigo = norm(row[0])
        desc = str(row[1] or "").strip() if len(row) > 1 else ""
        if codigo:
            omie.append((codigo, desc))
    print(f"Omie: {len(omie)} itens")

    conn = psycopg2.connect(url, cursor_factory=RealDictCursor, connect_timeout=60)
    cur = conn.cursor()
    cur.execute("SELECT ean, codigo_interno FROM produtos")
    rows = cur.fetchall()
    por_ean = {str(r["ean"]): (r["codigo_interno"] or "") for r in rows}
    por_limpo = {
        re.sub(r"[^0-9A-Za-z]", "", e).upper(): e for e in por_ean
    }
    print(f"Neon produtos: {len(por_ean)}")

    updates = []
    inserts = []
    for codigo, desc in omie:
        ean_alvo = codigo if codigo in por_ean else por_limpo.get(
            re.sub(r"[^0-9A-Za-z]", "", codigo).upper()
        )
        if ean_alvo:
            if (por_ean.get(ean_alvo) or "").strip() != codigo:
                updates.append((codigo, ean_alvo))
                por_ean[ean_alvo] = codigo
        else:
            marca, nome = "", desc
            partes = desc.split(" ", 1)
            if len(partes) > 1 and len(partes[0]) <= 20:
                marca, nome = partes[0], partes[1]
            inserts.append((codigo, nome or codigo, marca, codigo))
            por_ean[codigo] = codigo
            por_limpo[re.sub(r"[^0-9A-Za-z]", "", codigo).upper()] = codigo

    if updates:
        execute_batch(
            cur,
            "UPDATE produtos SET codigo_interno = %s, updated_at = CURRENT_TIMESTAMP WHERE ean = %s",
            updates,
            page_size=200,
        )
    if inserts:
        execute_batch(
            cur,
            "INSERT INTO produtos (ean, produto, marca, codigo_interno) VALUES (%s, %s, %s, %s) "
            "ON CONFLICT (ean) DO UPDATE SET codigo_interno = EXCLUDED.codigo_interno",
            inserts,
            page_size=200,
        )
    conn.commit()
    cur.execute("SELECT COUNT(*) AS n FROM produtos WHERE COALESCE(codigo_interno,'') <> ''")
    com_ci = cur.fetchone()["n"]
    cur.execute("SELECT COUNT(*) AS n FROM produtos")
    total = cur.fetchone()["n"]
    conn.close()
    print(f"OK updates={len(updates)} inserts={len(inserts)} | produtos={total} com_omie={com_ci}")


if __name__ == "__main__":
    main()
