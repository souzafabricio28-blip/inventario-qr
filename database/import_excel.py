import os
import re
from openpyxl import load_workbook
from .backend import atualizar_produto, criar_produto, listar_produtos


COLUNAS_MApeamento = {
    "ean": ["ean", "código de barras", "codigo", "código", "gtin", "barras"],
    "produto": ["produto", "descricao", "descrição", "descrição",
                "nome", "item", "descrição do produto", "descricao produto"],
    "marca": ["marca", "fabricante", "marca do produto"],
    "unidade_medida": ["unidade", "un", "unidade medida", "umed",
                       "unidade de medida", "und"],
    "base_especifica": ["base", "base específica", "base especifica",
                        "tipo base", "base técnica", "base tecnica"],
    "instrucao_dosagem": ["instrucao", "dosagem", "instrução",
                          "instrução dosagem", "regra", "proporção",
                          "instrucao dosagem"],
    "lote": ["lote", "batch", "lote fabricação", "lote fabricacao", "nro lote", "número lote"],
    "data_vencimento": ["data venci", "vencimento", "data venc", "validade", "dt venc",
                        "data de validade", "val", "venc"],
    "codigo_interno": ["codigo interno", "código interno", "id",
                        "cod_interno", "código", "referencia", "referência",
                        "sku", "código próprio", "codigo proprio"],
    "deposito": ["deposito", "depósito", "localização", "localizacao", "local", "armazém", "armazem"],
    "preco": ["preco", "preço", "valor", "preco venda", "preço venda",
              "preco custo", "preço custo", "valor unitário"],
}


def _normalizar_nome_coluna(nome):
    return re.sub(r"[_\-\s]", "", nome.strip().lower())


def _mapear_colunas(cabecalhos):
    mapping = {}
    cab_normalizados = {}
    for col in cabecalhos:
        cab_normalizados[col] = _normalizar_nome_coluna(col)

    for campo_alvo, possibilidades in COLUNAS_MApeamento.items():
        poss_norm = [_normalizar_nome_coluna(p) for p in possibilidades]
        for col_original, col_norm in cab_normalizados.items():
            if col_norm in poss_norm:
                mapping[campo_alvo] = col_original
                break

    return mapping


def importar_excel(caminho_arquivo):
    if not os.path.exists(caminho_arquivo):
        return False, f"Arquivo não encontrado: {caminho_arquivo}"

    try:
        wb = load_workbook(caminho_arquivo, read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            return False, "Planilha vazia"

        linhas = list(ws.iter_rows(values_only=True))
        if len(linhas) < 2:
            return False, "Planilha sem dados (apenas cabeçalho ou vazia)"

        cabecalhos = [str(c or "").strip() for c in linhas[0]]
        mapping = _mapear_colunas(cabecalhos)

        campos_obrigatorios = ["ean", "produto"]
        for campo in campos_obrigatorios:
            if campo not in mapping:
                return False, f"Coluna obrigatória '{campo}' não encontrada no Excel. Colunas encontradas: {', '.join(cabecalhos)}"

        importados = 0
        erros = []
        for idx, linha in enumerate(linhas[1:], start=2):
            if all(v is None or str(v).strip() == "" for v in linha):
                continue

            dados = {}
            for campo_alvo, col_orig in mapping.items():
                pos = cabecalhos.index(col_orig)
                val = linha[pos] if pos < len(linha) else None
                dados[campo_alvo] = str(val).strip() if val is not None else ""

            ean = dados.get("ean", "").replace(".0", "")
            if not ean or ean == "":
                continue

            produto_raw = dados.get("produto", "")
            marca = dados.get("marca", "")

            if not marca:
                partes = produto_raw.split(" ", 1)
                if len(partes) > 1 and len(partes[0]) <= 20:
                    marca = partes[0]
                    produto_raw = partes[1] if len(partes) > 1 else produto_raw

            unidade = dados.get("unidade_medida", "UN")
            base = dados.get("base_especifica", "")
            dosagem = dados.get("instrucao_dosagem", "")
            lote = dados.get("lote", "")
            data_venc = dados.get("data_vencimento", "")
            if data_venc and len(data_venc) >= 10:
                data_venc = data_venc[:10]
            codigo_interno = dados.get("codigo_interno", "")

            sucesso, msg = criar_produto(
                ean=ean,
                produto=produto_raw,
                marca=marca,
                unidade_medida=unidade,
                base_especifica=base,
                instrucao_dosagem=dosagem,
                lote=lote,
                data_vencimento=data_venc,
                codigo_interno=codigo_interno,
            )
            if sucesso:
                importados += 1
            else:
                if "UNIQUE" not in msg:
                    erros.append(f"Linha {idx} (EAN {ean}): {msg}")

        wb.close()
        resumo = f"{importados} produtos importados"
        if erros:
            resumo += f", {len(erros)} erros ignorados"

        return True, resumo

    except Exception as e:
        return False, f"Erro ao importar: {str(e)}"


def _norm_codigo_omie(val):
    """Normaliza código Omie (ex.: 29912.0 → 29912)."""
    if val is None:
        return ""
    c = str(val).strip()
    if c.endswith(".0") and c[:-2].replace("-", "").replace(".", "").isdigit():
        c = c[:-2]
    return c


def caminho_lista_rt_oficial():
    """Compat: caminhos da lista RT oficial."""
    from .rt_lista import caminho_lista_rt
    p = caminho_lista_rt()
    return [p] if p else []



def sincronizar_rt_oficial(caminho_arquivo=None, limpar_antigos=False):
    """
    Espelha a planilha RT oficial no banco operacional.
    A planilha é a fonte única de catálogo (ean + codigo_omie).
    limpar_antigos=True remove do banco produtos cujo ean não está na planilha.
    """
    from .backend import (
        criar_produto as _criar,
        atualizar_produto as _atualizar,
        vincular_codigo,
        excluir_produto_completo,
        _buscar_produto_db,
        _listar_produtos_db,
    )
    from .rt_lista import carregar_lista_rt, chaves_oficiais, caminho_lista_rt

    if not caminho_arquivo:
        caminho_arquivo = caminho_lista_rt()
    if not caminho_arquivo or not os.path.exists(caminho_arquivo):
        return False, "Arquivo lista_produtos_RToficial.xlsx não encontrado"

    # Copiar para data/ se veio de Documents (Vercel/local canônico)
    raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    dest = os.path.join(raiz, "data", "lista_produtos_RToficial.xlsx")
    try:
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        if os.path.abspath(caminho_arquivo) != os.path.abspath(dest):
            import shutil
            shutil.copy2(caminho_arquivo, dest)
            caminho_arquivo = dest
    except Exception:
        pass

    ok_load, msg_load = carregar_lista_rt(force=True)
    if not ok_load:
        return False, msg_load

    try:
        wb = load_workbook(caminho_arquivo, read_only=True, data_only=True)
        ws = wb["Lista Unica"] if "Lista Unica" in wb.sheetnames else wb.active
        linhas = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception as e:
        return False, f"Erro ao ler lista RT: {e}"

    if len(linhas) < 2:
        return False, "Planilha RT vazia"

    cabecalhos = [str(c or "").strip().lower() for c in linhas[0]]

    def _idx(*nomes):
        for n in nomes:
            if n in cabecalhos:
                return cabecalhos.index(n)
        return None

    i_ean = _idx("ean", "codigo_barras", "gtin")
    i_cod = _idx("codigo")
    i_desc = _idx("descricao", "produto", "nome")
    i_omie = _idx("codigo_omie")

    if i_ean is None and i_cod is None and i_omie is None:
        return False, f"Colunas ean/codigo_omie não encontradas: {cabecalhos}"
    if i_desc is None:
        return False, f"Coluna descricao não encontrada: {cabecalhos}"

    def _cell(linha, idx):
        if idx is None or idx >= len(linha):
            return ""
        val = linha[idx]
        if idx in (i_ean, i_cod, i_omie):
            return _norm_codigo_omie(val)
        return str(val or "").strip()

    criados = 0
    atualizados = 0
    aliases = 0
    erros = 0

    for linha in linhas[1:]:
        if not linha or all(v is None or str(v).strip() == "" for v in linha):
            continue
        codigo = _cell(linha, i_cod)
        omie = _cell(linha, i_omie) or codigo
        ean = _cell(linha, i_ean) or omie or codigo
        if not ean:
            continue
        descricao = _cell(linha, i_desc) or ean
        codigo_interno = omie or codigo or ean

        marca = ""
        nome = descricao
        partes = descricao.split(" ", 1)
        if len(partes) > 1 and len(partes[0]) <= 20:
            marca, nome = partes[0], partes[1]

        # PK oficial = EAN da planilha (não o codigo_omie)
        existente = _buscar_produto_db(ean)
        corrigido_manual = bool(existente.get("editado_manual")) if existente else False
        if existente and (existente.get("ean") or "") == ean:
            ean_alvo = ean
            mudou = (
                (existente.get("produto") or "") != nome
                or (existente.get("marca") or "") != marca
                or (existente.get("codigo_interno") or "") != codigo_interno
            )
            if mudou and not corrigido_manual:
                _atualizar(
                    ean_alvo,
                    produto=nome,
                    marca=marca,
                    codigo_interno=codigo_interno,
                )
                atualizados += 1
        else:
            ok, msg = _criar(
                ean=ean, produto=nome, marca=marca, codigo_interno=codigo_interno,
                editado_manual=False,
            )
            if ok:
                criados += 1
                ean_alvo = ean
            elif "UNIQUE" in (msg or "").upper():
                _atualizar(ean, produto=nome, marca=marca, codigo_interno=codigo_interno)
                atualizados += 1
                ean_alvo = ean
            else:
                erros += 1
                continue

        for alias in (codigo, omie, ean):
            if alias and alias != ean_alvo:
                ok_v, _ = vincular_codigo(alias, ean_alvo, origem="rt_oficial", sobrescrever=False)
                if ok_v:
                    aliases += 1

    removidos = 0
    if limpar_antigos:
        eans_ok, _omies_ok = chaves_oficiais()
        for p in _listar_produtos_db(""):
            ean_p = (p.get("ean") or "").strip()
            # Mantém o que tem EAN igual ao da planilha (fonte oficial)
            if ean_p and ean_p in eans_ok:
                continue
            # Produtos com correção manual não são apagados pela planilha
            if bool(p.get("editado_manual")):
                continue
            try:
                excluir_produto_completo(ean_p)
                removidos += 1
            except Exception:
                erros += 1

    carregar_lista_rt(force=True)
    return True, (
        f"Catálogo RT oficial aplicado ({os.path.basename(caminho_arquivo)}): "
        f"{criados} novos, {atualizados} atualizados, {aliases} vínculos"
        + (f", {removidos} antigos removidos" if removidos else "")
        + (f", {erros} erros" if erros else "")
    )


def sincronizar_omie(caminho_arquivo=None):
    """
    Cruza a planilha Omie (codigo + descrição) com a tabela produtos:
    - preenche codigo_interno com o código Omie
    - cadastra produtos Omie que ainda não existem (ean = código)
    """
    if not caminho_arquivo:
        raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        candidatos = [
            os.path.join(raiz, "omie_produtos.xlsx"),
            os.path.join(os.path.expanduser("~"), "Documents",
                         "lista de produtos cadastrados no omie.xlsx"),
        ]
        caminho_arquivo = next((p for p in candidatos if os.path.exists(p)), None)

    if not caminho_arquivo or not os.path.exists(caminho_arquivo):
        return False, "Arquivo Omie não encontrado (omie_produtos.xlsx)"

    try:
        wb = load_workbook(caminho_arquivo, read_only=True, data_only=True)
        ws = wb.active
        linhas = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception as e:
        return False, f"Erro ao ler Omie: {e}"

    if len(linhas) < 2:
        return False, "Planilha Omie vazia"

    with_prods = listar_produtos()
    por_ean = {}
    por_limpo = {}
    for p in with_prods:
        ean = str(p["ean"])
        ci = (p.get("codigo_interno") or "")
        por_ean[ean] = ci
        por_limpo[re.sub(r"[^0-9A-Za-z]", "", ean).upper()] = ean

    atualizados = 0
    criados = 0
    for linha in linhas[1:]:
        if not linha or (linha[0] is None and (len(linha) < 2 or not linha[1])):
            continue
        codigo = _norm_codigo_omie(linha[0])
        descricao = str(linha[1] or "").strip() if len(linha) > 1 else ""
        if not codigo:
            continue

        ean_alvo = codigo if codigo in por_ean else por_limpo.get(
            re.sub(r"[^0-9A-Za-z]", "", codigo).upper()
        )

        if ean_alvo:
            prod_atual = None
            if (por_ean.get(ean_alvo) or "").strip() != codigo:
                try:
                    from .backend import _buscar_produto_db as _buscar
                    prod_atual = _buscar(ean_alvo)
                except Exception:
                    prod_atual = None
                # Não sobrescreve o códigos dos produtos corrigidos na mão
                if not (bool(prod_atual.get("editado_manual")) if prod_atual else False):
                    atualizar_produto(ean_alvo, codigo_interno=codigo)
                    por_ean[ean_alvo] = codigo
                    atualizados += 1
        else:
            marca = ""
            nome = descricao
            partes = descricao.split(" ", 1)
            if len(partes) > 1 and len(partes[0]) <= 20:
                marca, nome = partes[0], partes[1]
            ok, _ = criar_produto(
                ean=codigo, produto=nome or codigo, marca=marca,
                codigo_interno=codigo, editado_manual=False,
            )
            if ok:
                por_ean[codigo] = codigo
                por_limpo[re.sub(r"[^0-9A-Za-z]", "", codigo).upper()] = codigo
                criados += 1

    return True, f"Omie sincronizado: {atualizados} atualizados, {criados} novos"
